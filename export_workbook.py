import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "interview-prep.xlsx"


def read_text(name):
    return (ROOT / name).read_text(encoding="utf-8")


def markdown_table_rows(text):
    rows = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped.startswith("|") or not stripped.endswith("|"):
            continue
        cells = [cell.strip() for cell in stripped[1:-1].split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    return rows


def add_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([str(value) if value is not None else "" for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174A5B")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        max_length = max((len(str(cell.value or "")) for cell in column), default=10)
        sheet.column_dimensions[letter].width = min(max(max_length * 0.8, 14), 48)
    return sheet


workbook = Workbook()
overview = workbook.active
overview.title = "Overview"
overview_rows = [
    ["Candidate", "Diksha Parganiha"],
    ["Target roles", "Senior Backend Engineer; Senior Platform Engineer; Lead Engineer; Staff-track roles"],
    ["Research date", "2026-10-19"],
    ["Source boundary", "Public company guidance, primary technical documentation, practitioner guidance, and surveys; not a complete census of interviews."],
    ["Resume status", "No resume/CV was present in the workspace. Resume-derived question generation and individualized STAR answers require the resume."],
    ["Company data warning", "Ranking, role match, and difficulty are heuristic. Rounds, compensation, Go/Kafka use, and AI emphasis require validation against the exact role and location."],
    ["Read first", "Market Research; Topic Roadmap; Execution Plan; Behavioral Bank; Knowledge Checks."],
]
for row in overview_rows:
    overview.append(row)
overview.column_dimensions["A"].width = 26
overview.column_dimensions["B"].width = 110
for row in overview.iter_rows():
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if cell.column == 1:
            cell.font = Font(bold=True, color="174A5B")

company_rows = list(csv.reader((ROOT / "company-targets.csv").open(newline="", encoding="utf-8")))
add_sheet(workbook, "Company Targets", company_rows[0], company_rows[1:])

roadmap_text = read_text("topic-roadmap.md")
roadmap_section = roadmap_text.split("## Topic Catalog", 1)[1].split("## Go Master Roadmap", 1)[0]
roadmap_rows = markdown_table_rows(roadmap_section)
add_sheet(workbook, "Topic Roadmap", roadmap_rows[0], roadmap_rows[1:])

market_text = read_text("market-research.md")
market_section = market_text.split("## Market Signal Register", 1)[1].split("## Increasing, Stable", 1)[0]
market_rows = markdown_table_rows(market_section)
add_sheet(workbook, "Market Findings", market_rows[0], market_rows[1:])

behavior_text = read_text("behavioral-question-bank.md")
behavior_rows = []
category = ""
for line in behavior_text.splitlines():
    heading = re.match(r"## (\d+)\. (.+)", line)
    if heading:
        category = heading.group(2)
        continue
    question = re.match(r"(\d+)\. (.+)", line)
    if question:
        behavior_rows.append([int(question.group(1)), category, question.group(2), "STAR-L: Situation, Task, Actions, Result, Learning", "Use a true example; see bank scorecard. Individualized answers require resume evidence."])
add_sheet(workbook, "Behavioral Bank", ["Number", "Competency", "Question", "Answer Structure", "Personalization Note"], behavior_rows)

plan_text = read_text("execution-plan.md")
plan_rows = []
section = ""
for line in plan_text.splitlines():
    heading = re.match(r"## (.+)", line)
    if heading:
        section = heading.group(1)
        continue
    if line.strip().startswith("|") and line.strip().endswith("|"):
        cells = [cell.strip() for cell in line.strip()[1:-1].split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        plan_rows.append([section] + cells)
add_sheet(workbook, "Execution Plan", ["Plan Section", "Period", "Focus", "Required Outputs"], plan_rows)

checks_text = read_text("knowledge-checks.md")
checks_section = checks_text.split("## Four-Level Check Set", 1)[1].split("## Answer Key", 1)[0]
checks_rows = markdown_table_rows(checks_section)
add_sheet(workbook, "Knowledge Checks", checks_rows[0], checks_rows[1:])

resources_text = roadmap_text.split("## Free Resource Index", 1)[1]
resource_rows = markdown_table_rows(resources_text)
add_sheet(workbook, "Free Resources", resource_rows[0], resource_rows[1:])

for sheet in workbook.worksheets:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

workbook.save(OUTPUT)
print(f"Created {OUTPUT}")
print(f"Companies: {len(company_rows) - 1}; behavioral questions: {len(behavior_rows)}; roadmap topics: {len(roadmap_rows) - 1}")
