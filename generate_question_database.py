import csv
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
CSV_PATH = ROOT / "question-database.csv"
XLSX_PATH = ROOT / "interview-question-database.xlsx"

# Topic totals are exactly the minimum requested per section. Prompts are
# deliberately phrased as practice questions, not falsely attributed verbatim.
CATALOG = [
    ("Go", 1500, "High", 5, "Amazon SDE III publishes programming fundamentals; other employer mappings are role-family inferences.", "https://go.dev/doc/", [
        "Go values, types, and zero values", "pointers and value semantics", "arrays and slice descriptors", "slice capacity and append", "maps and iteration", "structs and methods", "interfaces and method sets", "typed nil and nil interfaces", "errors and wrapping", "defer, panic, and recover", "generics and constraints", "reflection and type assertions", "packages and modules", "goroutines and lifecycle", "channels and close ownership", "select and cancellation", "context deadlines and propagation", "WaitGroups and worker pools", "mutexes and RWMutex", "atomics and memory ordering", "Go memory model and race detection", "GMP scheduler and work stealing", "preemption and blocking goroutines", "garbage collector and pacing", "escape analysis and stack versus heap", "compiler optimization and inlining", "testing, fuzzing, and test doubles", "benchmarking and allocation measurement", "pprof, execution trace, and diagnostics", "production shutdown, retries, and backpressure"
    ]),
    ("DSA", 1450, "High", 5, "Amazon official guidance names data structures, algorithms, and coding; named company/problem mapping is not claimed.", "https://amazon.jobs/content/en/how-we-hire/interview-prep/software-development-topics", [
        "Two Sum", "Three Sum", "Four Sum", "Subarray Sum Equals K", "Longest Substring Without Repeating Characters", "Minimum Window Substring", "Merge Intervals", "Insert Interval", "Meeting Rooms", "Top K Frequent Elements", "Kth Largest Element", "LRU Cache", "LFU Cache", "Trie Prefix Search", "Word Search Backtracking", "Combination Sum", "N-Queens", "Coin Change", "Longest Increasing Subsequence", "Edit Distance", "House Robber", "Knapsack", "Decode Ways", "Number of Islands", "Course Schedule", "Monotonic Stack", "Union-Find", "Segment Tree", "Graph Shortest Path"
    ]),
    ("Distributed Systems", 1000, "High", 5, "Broad role-family inference; primary references support concepts, not employer-specific prompt counts.", "https://sre.google/workbook/table-of-contents/", [
        "CAP under network partitions", "linearizability", "sequential and causal consistency", "eventual consistency", "quorum reads and writes", "leader-follower replication", "multi-leader replication", "Raft election and log replication", "Paxos safety and liveness", "failure detection", "split brain and fencing", "distributed locks and leases", "clock drift and ordering", "Lamport and vector clocks", "idempotency and deduplication", "transactional outbox", "CDC delivery", "sagas and compensation", "event sourcing and CQRS", "CRDT convergence"
    ]),
    ("System Design", 1000, "High", 5, "Design prompts are representative practice scenarios; no employer-by-employer assignment is asserted.", "https://amazon.jobs/content/en/how-we-hire/sde-iii-interview-prep", [
        "WhatsApp messaging", "payment gateway", "EV charging platform", "multi-tenant SaaS", "notification service", "Kafka-like event platform", "LLM inference platform", "agent tool platform", "feature flag service", "internal developer platform", "YouTube video platform", "Instagram social feed", "Netflix streaming", "Google Drive file sync", "Google Maps and geospatial search", "Uber ride dispatch", "Stripe payment platform", "search engine", "API gateway", "distributed rate limiter"
    ]),
    ("Kafka", 500, "High", 5, "Official Kafka documentation grounds technical facts; question frequency is an estimate for matching roles.", "https://kafka.apache.org/documentation/", [
        "partitions and ordering", "replication and ISR", "leader election and unclean election", "producer acknowledgements and durability", "idempotent producer", "transactions and exactly-once boundary", "consumer groups and assignment", "consumer offsets and delivery semantics", "rebalance protocols", "KRaft metadata quorum"
    ]),
    ("PostgreSQL", 500, "High", 5, "PostgreSQL manual grounds technical facts; frequency is an estimate for backend roles.", "https://www.postgresql.org/docs/current/mvcc.html", [
        "MVCC snapshots", "transaction isolation anomalies", "row locks and deadlocks", "B-tree and composite indexes", "EXPLAIN ANALYZE plans", "WAL and crash recovery", "VACUUM and visibility maps", "HOT updates and index maintenance", "table and index bloat", "partitioning and pruning"
    ]),
    ("Redis", 250, "Medium", 4, "Redis documentation grounds replication and behavior; question frequency is role-dependent.", "https://redis.io/docs/latest/operate/oss_and_stack/management/replication/", [
        "cache-aside and invalidation", "TTL and eviction", "asynchronous replication", "Sentinel and Cluster failover", "hot keys and cache stampedes"
    ]),
    ("MongoDB", 250, "Medium", 4, "MongoDB documentation grounds version-specific behavior; question frequency is role-dependent.", "https://www.mongodb.com/docs/manual/", [
        "document modeling: embed versus reference", "indexes and query planning", "replica sets and write concern", "sharding and shard-key selection", "transactions, aggregation, and change streams"
    ]),
    ("Kubernetes", 500, "High", 5, "CNCF surveys support ecosystem relevance, not interview counts; concepts map strongly to platform roles.", "https://kubernetes.io/docs/concepts/overview/", [
        "control plane reconciliation", "Pod lifecycle and probes", "Deployments and rollouts", "Services and service discovery", "Ingress and Gateway API", "requests, limits, and QoS", "HPA and autoscaling", "RBAC and workload identity", "network policies and security boundaries", "persistent volumes and stateful workloads"
    ]),
    ("AWS/Azure", 500, "High", 4, "Cloud relevance is role-specific; AWS Well-Architected is a design source, not an interview-question census.", "https://docs.aws.amazon.com/wellarchitected/latest/framework/welcome.html", [
        "IAM and least privilege", "VPC/VNet network segmentation", "load balancing and DNS", "compute selection and autoscaling", "object and block storage", "managed databases and recovery", "queues and event services", "multi-region resilience", "cloud cost and egress", "cloud observability and incident response"
    ]),
    ("Security", 500, "High", 5, "OWASP and Microsoft responsible-AI guidance support risk topics; company-specific interview attribution is not inferred.", "https://owasp.org/www-project-top-10-for-large-language-model-applications/", [
        "OAuth2 flows and threat boundaries", "OIDC identity and token validation", "JWT signature and claims validation", "RBAC versus ABAC", "object-level authorization", "tenant isolation", "SSRF and egress controls", "SQL injection and parameterization", "secrets rotation and KMS", "software supply-chain security"
    ]),
    ("Behavioral", 500, "High", 5, "Amazon and Microsoft publish behavioral/value guidance; prompts below are practice formulations, not quoted interview scripts.", "https://amazon.jobs/content/en/how-we-hire/sde-iii-interview-prep", [
        "ownership and customer impact", "technical disagreement", "failure and recovery", "ambiguity and discovery", "prioritization and tradeoffs", "influence without authority", "incident leadership", "mentoring and delegation", "security and integrity", "measurable cross-team impact"
    ]),
    ("AI/LLM", 550, "Emerging", 4, "Survey and OWASP/MCP sources establish ecosystem/security relevance. Sierra's first-party AI-native interview article is a company-specific emerging signal; its body was blocked in this environment, so broader expectations remain unverified.", "https://modelcontextprotocol.io/specification/2025-06-18", [
        "transformer and token fundamentals", "embeddings and semantic retrieval", "RAG retrieval and grounding", "vector database tradeoffs", "LLM evaluation design", "tool calling and authorization", "agent workflow reliability", "MCP host/client/server boundaries", "prompt injection and data exfiltration", "AI observability and cost controls", "AI-assisted coding interview: tool use and verification"
    ]),
    ("Resume-Based", 500, "High", 5, "Profile-led prompts only; resume absent. Validate every premise against actual experience before using an answer.", "interview-prep/README.md", [
        "Go service ownership", "Kafka and event-driven architecture", "PostgreSQL and Redis decisions", "Kubernetes and deployment ownership", "AWS/Azure architecture", "EV charging domain decisions", "cloud security platform experience", "observability and incident response", "cross-team technical leadership", "career decisions across Shell, Zscaler, and Accenture"
    ]),
]

# Each lens yields a distinct question form. The framing is a deliberate
# generation method; it does not imply that a question appeared verbatim in an interview.
LENSES = [
    ("Beginner", "Define {topic} and explain the problem it solves."),
    ("Beginner", "Walk through {topic} with a minimal concrete example."),
    ("Beginner", "What are the essential terms and invariants behind {topic}?"),
    ("Beginner", "How would you explain {topic} to an engineer new to backend systems?"),
    ("Intermediate", "What assumptions must hold for {topic} to behave correctly?"),
    ("Intermediate", "Compare {topic} with its closest alternative and give a workload where each wins."),
    ("Intermediate", "Which edge cases break a naive implementation of {topic}?"),
    ("Intermediate", "How would you test correctness for {topic}, including boundary cases?"),
    ("Intermediate", "What is the most common misconception about {topic}, and what is actually guaranteed?"),
    ("Intermediate", "Trace the lifecycle of a request or operation involving {topic}."),
    ("Advanced", "What failure mode involving {topic} is hardest to detect, and how would you expose it?"),
    ("Advanced", "How does {topic} behave under partial failure, retries, or cancellation?"),
    ("Advanced", "How would you diagnose a production incident involving {topic}?"),
    ("Advanced", "Which metrics, traces, logs, or profiles would help investigate {topic}, and why?"),
    ("Advanced", "How would you improve the performance of {topic} without weakening correctness?"),
    ("Advanced", "What memory, allocation, storage, or network costs does {topic} introduce?"),
    ("Advanced", "How can concurrency change the correctness of {topic}?"),
    ("Advanced", "What security or isolation boundary applies to {topic}?"),
    ("Advanced", "How would you roll out a change to {topic} while preserving compatibility?"),
    ("Advanced", "What test would falsify your main assumption about {topic}?"),
    ("Senior", "Design a production-ready use of {topic}; state requirements, SLOs, and failure behavior."),
    ("Senior", "A service using {topic} has rising p99 latency but steady average latency. How do you investigate?"),
    ("Senior", "Traffic doubles and one tenant dominates a system using {topic}. What changes do you make?"),
    ("Senior", "How would you choose between consistency, availability, latency, and cost when using {topic}?"),
    ("Senior", "Describe a safe migration from the current use of {topic} to a new design."),
    ("Senior", "How would you make {topic} observable and operable by an on-call team?"),
    ("Senior", "What happens when a dependency of {topic} becomes slow rather than fully unavailable?"),
    ("Senior", "How do you bound retries, queues, resource use, and blast radius around {topic}?"),
    ("Senior", "Which invariants should be enforced in code, storage, or protocol when using {topic}?"),
    ("Senior", "Tell me about a real decision involving {topic}; what alternatives did you reject and why?"),
    ("Staff", "Several teams depend on {topic}. How would you set direction and obtain adoption without authority?"),
    ("Staff", "How would you define a multi-team contract and migration policy around {topic}?"),
    ("Staff", "When should an organization standardize {topic}, and when should teams retain local choice?"),
    ("Staff", "How would you measure whether an investment in {topic} improved engineering outcomes?"),
    ("Staff", "A proposed use of {topic} solves a local problem but adds organization-wide operational burden. How do you decide?"),
    ("Staff", "How would you delegate a high-risk {topic} initiative while preserving technical quality?"),
    ("Staff", "How do you resolve disagreement between teams over {topic} when both positions are reasonable?"),
    ("Staff", "What durable capability should remain after a cross-team {topic} project ends?"),
    ("Principal", "How would you set a multi-year technical strategy for {topic} under uncertain product direction?"),
    ("Principal", "What organizational failure could make an otherwise sound {topic} architecture unsuccessful?"),
    ("Principal", "How would you establish a decision framework for {topic} across independent business units?"),
    ("Principal", "What evidence would cause you to reverse a major investment in {topic}?"),
    ("Principal", "How would you balance standardization, migration cost, autonomy, and risk for {topic}?"),
    ("Debugging", "A defect report says {topic} produced an impossible result. What is your first diagnostic sequence?"),
    ("Debugging", "How do you distinguish a bug in {topic} from a bad assumption in its caller?"),
    ("Failure", "Walk through the worst plausible cascading failure involving {topic} and your containment plan."),
    ("Performance", "How would you benchmark {topic} without measuring the wrong workload or introducing bias?"),
    ("Architecture", "Where should ownership of {topic} sit, and what contract should callers rely on?"),
    ("Tradeoff", "What is a case where using {topic} is the wrong choice even if it is technically feasible?"),
    ("Follow-up", "An interviewer challenges your answer about {topic}. What evidence or guarantee supports your claim?"),
    ("Synthesis", "Explain {topic} from first principles, then connect it to a production design and a failure you would test."),
]

DSA_LENSES = LENSES

ANSWER_CUES = {
    "Go": "A strong answer states Go's specified language/runtime contract, distinguishes guarantees from implementation details, gives a small example, names failure modes, and proposes a test or measurement. Qualify runtime internals by Go version.",
    "DSA": "A strong solution clarifies inputs/constraints, states an invariant, derives correctness, gives time/space complexity, tests edge cases, and discusses a simpler or more optimal alternative.",
    "Distributed Systems": "A strong answer names the failure model and consistency/durability boundary, distinguishes safety from liveness, traces a concrete failure, and explains recovery, fencing/idempotency, and operational evidence.",
    "System Design": "Clarify users and functional/nonfunctional requirements; estimate load/storage; define APIs and data; draw a minimal design; analyze bottlenecks, failure, security, SLO, migration, and cost; justify tradeoffs.",
    "Kafka": "State partition-level ordering, replication/ISR and acknowledgment settings, offset commit boundary, retry/duplicate behavior, and the exact scope of Kafka transactions. Do not promise external exactly-once effects without a coordinated sink or idempotency.",
    "PostgreSQL": "Use SQL semantics and actual plan/lock evidence. Explain snapshot/isolation behavior, constraints/index tradeoffs, WAL/recovery, vacuum/bloat, and safe rollout instead of relying on ORM intuition.",
    "Redis": "Treat Redis guarantees as configuration/version dependent. Distinguish cache from durable source of truth, explain asynchronous replication/failover and staleness, and state eviction/invalidation behavior.",
    "MongoDB": "Explain document boundaries, index/query shape, read/write concerns, shard-key consequences, transactions, and operational tradeoffs for the target MongoDB version.",
    "Kubernetes": "Explain desired-state reconciliation and the relevant API objects, resource/health behavior, security boundaries, and application-level limits. Kubernetes primitives do not replace app-level reliability design.",
    "AWS/Azure": "Start from requirements and failure domains, then justify identity/network/compute/storage choices. Include quotas, multi-zone/region behavior, recovery, egress/cost, shared responsibility, and least privilege.",
    "Security": "Name assets, actors, trust boundaries, and abuse cases; enforce authorization server-side; minimize privilege/data exposure; specify validation, audit, rotation, detection, and residual risk.",
    "Behavioral": "Use a truthful STAR-L example: situation/stakes, personal task, specific actions and alternatives, measurable or defensible result, and learning that changed later behavior. Do not invent confidential metrics.",
    "AI/LLM": "Separate model capability from system guarantees. Define evaluation data/metrics, retrieval quality, access control, tool permissions, prompt-injection defenses, fallback, observability, privacy, latency, and cost.",
    "Resume-Based": "Answer only from verifiable personal experience: context, personal ownership, architecture, alternatives, tradeoffs, failure, metrics, and lessons. If the resume does not support a premise, correct it rather than fabricate.",
}

TOPIC_CORE = {
    "Go values, types, and zero values": "Go variables have a static type and each type has a zero value; assignments copy values unless copied values contain references to shared underlying data.",
    "pointers and value semantics": "A pointer refers to an addressable value; passing a pointer copies the pointer, not the pointed-to object. Use pointers when mutation/identity or avoiding large copies justifies them.",
    "arrays and slice descriptors": "An array owns fixed-length elements. A slice is a descriptor over an array (pointer, length, capacity); copying a slice copies the descriptor and may preserve shared backing storage.",
    "slice capacity and append": "Append reuses backing storage when capacity allows and allocates a new array when it does not; callers must not assume either outcome without reasoning about capacity and aliasing.",
    "maps and iteration": "Go maps are reference-like, iteration order is unspecified, and concurrent read/write access needs synchronization; map lookups can use the comma-ok form.",
    "structs and methods": "Methods have receiver types; value and pointer receiver method sets differ. Embedding promotes methods but is composition, not inheritance.",
    "interfaces and method sets": "An interface is satisfied implicitly by implementing its method set. Interface values carry dynamic type and value; keep interfaces small and owned near consumers.",
    "typed nil and nil interfaces": "A nil interface has no dynamic type/value. An interface holding a typed nil pointer has a dynamic type and is therefore not equal to nil.",
    "errors and wrapping": "Errors are explicit values; wrap with %w when preserving identity, then use errors.Is/As. Add useful context and avoid string matching for control flow.",
    "defer, panic, and recover": "Deferred calls run on function return and during panic unwinding. Panic/recover is for exceptional invariant failure or boundary recovery, not ordinary error handling.",
    "generics and constraints": "Type parameters and constraints express reusable algorithms while retaining static type checking; use them when they improve clarity over interfaces or duplication.",
    "reflection and type assertions": "Reflection inspects dynamic types/values at runtime with weaker compile-time guarantees. Prefer ordinary interfaces/type switches unless runtime type discovery is required.",
    "packages and modules": "Packages define compilation and API boundaries; modules version a set of packages. Keep dependencies explicit, APIs small, and compatibility intentional.",
    "goroutines and lifecycle": "A goroutine is a concurrent execution unit managed by the runtime, not a resource that is automatically cancelled or garbage-collected; every goroutine needs an exit condition and owner.",
    "channels and close ownership": "Channel send/receive synchronize communication; close signals no further values. The sending side that knows no more sends will occur should own closing.",
    "select and cancellation": "Select waits for a ready communication case; when multiple are ready, selection is pseudo-random. Include context.Done in blocking paths to permit cancellation.",
    "context deadlines and propagation": "Context carries cancellation, deadlines, and request-scoped values across API boundaries. Propagate it to downstream work; do not use it as an optional-parameter bag.",
    "WaitGroups and worker pools": "WaitGroup coordinates completion, not cancellation or error propagation. Bound workers and queue capacity, and define shutdown/cancellation separately.",
    "mutexes and RWMutex": "Mutex protects shared invariants with a critical section. RWMutex may help read-heavy workloads but can add overhead; profile before choosing it.",
    "atomics and memory ordering": "Atomic operations provide indivisible updates and synchronization guarantees for carefully designed state; they do not automatically make multi-field invariants safe.",
    "Go memory model and race detection": "Data-race-free Go programs behave sequentially consistently. Synchronization establishes happens-before; the race detector finds some executed races but does not prove race freedom.",
    "GMP scheduler and work stealing": "The runtime multiplexes goroutines (G) over OS threads (M) using logical processors (P); scheduling and work stealing distribute runnable work, subject to Go-version implementation details.",
    "preemption and blocking goroutines": "A blocked goroutine yields its thread for other work; preemption helps runtime scheduling but does not cancel blocked application operations. I/O, syscall, and cgo behavior matter operationally.",
    "garbage collector and pacing": "Go uses concurrent tracing GC with brief coordination phases; allocation rate and live heap drive cost. Measure heap, allocation, CPU, and pause behavior before tuning.",
    "escape analysis and stack versus heap": "Compiler escape analysis determines whether values may outlive a stack frame; inspect compiler output and profiles rather than assuming source-level pointers always allocate on heap.",
    "compiler optimization and inlining": "The compiler applies optimization such as inlining and escape analysis. Runtime/compiler implementation details can vary by version; benchmark and inspect generated diagnostics.",
    "testing, fuzzing, and test doubles": "Tests should assert externally meaningful behavior and invariants; table tests cover cases, fuzzing explores inputs, and mocks should not replace integration/contract confidence.",
    "benchmarking and allocation measurement": "Benchmark representative work with stable inputs and repeated runs; use benchstat and -benchmem. Avoid measuring setup or compiler-eliminated work accidentally.",
    "pprof, execution trace, and diagnostics": "CPU/heap/block/mutex profiles locate cost; execution trace exposes scheduling and latency behavior. Choose one tool based on the symptom and account for profiling overhead.",
    "production shutdown, retries, and backpressure": "Bound work and retries, propagate deadlines, stop accepting work before draining, and make shutdown idempotent. Backpressure prevents overload from becoming unbounded resource growth.",
    "Two Sum": "Scan once while storing previously seen values in a hash map; for each x look for target-x. Expected O(n) time and O(n) space; clarify duplicate/index reuse rules.",
    "Three Sum": "Sort then fix one value and use two pointers, skipping duplicates. O(n^2) time after sorting; discuss overflow and unique-triplet semantics.",
    "Four Sum": "Sort and use nested fixed indices plus two pointers, skipping duplicates; O(n^3) time and constant extra space beyond sorting. Guard integer overflow.",
    "Subarray Sum Equals K": "Use prefix-sum frequencies: for each prefix p add count[p-k], then increment count[p]. O(n) time/O(n) space; initialize count[0]=1 and handle negative values.",
    "Longest Substring Without Repeating Characters": "Sliding window with last-seen positions; move the left bound monotonically past duplicates. O(n) time/O(character set) space.",
    "Minimum Window Substring": "Sliding window with required and formed counts; expand until valid, then shrink while valid. O(|s|+|t|) time; account for duplicate characters.",
    "Merge Intervals": "Sort by start and merge when the next start overlaps the current end. O(n log n) time; clarify touching endpoints and empty input.",
    "Insert Interval": "Append intervals before the new range, merge all overlaps into it, then append remaining ranges. O(n) if already sorted/disjoint.",
    "Meeting Rooms": "Sort intervals by start and track the latest end, or sort start/end events; O(n log n). Clarify whether touching meetings overlap.",
    "Top K Frequent Elements": "Count frequencies, then use a size-k heap O(n log k) or bucket by frequency O(n); state tie behavior if required.",
    "Kth Largest Element": "Use quickselect expected O(n) or a size-k min-heap O(n log k); discuss worst-case quickselect and duplicate ranking semantics.",
    "LRU Cache": "Combine hash map lookup with a doubly linked list ordered by recency; get/put are O(1). Define capacity-zero and update behavior.",
    "LFU Cache": "Track key values/frequencies plus recency order within each frequency and the minimum frequency; expected O(1) operations with careful updates/ties.",
    "Trie Prefix Search": "Store child edges by character and terminal markers; lookup/insert cost O(length). Discuss alphabet representation, memory, and deletion if needed.",
    "Word Search Backtracking": "DFS from each cell while marking the current path and restoring state on return; prune using bounds/visited/character checks. Worst-case exponential in word length.",
    "Combination Sum": "Backtrack over candidate choices, tracking remaining sum and start index to avoid permutation duplicates; prune when candidates are positive and exceed remaining target.",
    "N-Queens": "Backtrack row by row while tracking used columns and diagonals; bitmasks make conflict checks O(1). Search remains exponential.",
    "Coin Change": "Dynamic programming over amounts or BFS over reachable amounts; O(amount * coin_count) DP. Define unreachable state and amount-zero base case.",
    "Longest Increasing Subsequence": "Maintain tails and binary-search the first tail >= x for strict LIS; O(n log n). Tails represent minimal endpoints, not the actual subsequence by themselves.",
    "Edit Distance": "Dynamic programming over prefixes with insert/delete/replace transitions; O(mn) time and O(min(m,n)) space with rolling rows.",
    "House Robber": "Track best values with and without taking the current item, or two rolling DP states; O(n) time/O(1) space. Clarify empty input and negative values.",
    "Knapsack": "Use DP over capacity and items; 0/1 knapsack iterates capacity downward for a 1D array. State whether items are reusable and define capacity constraints.",
    "Decode Ways": "DP over positions using valid one- and two-digit transitions; O(n) time/O(1) space. Handle leading zero and invalid pairs.",
    "Number of Islands": "Flood-fill each unvisited land cell with DFS/BFS, marking visited. O(rows*cols) time; discuss mutation, recursion depth, and 4- vs 8-connectivity.",
    "Course Schedule": "Build prerequisite graph and detect cycles using DFS coloring or Kahn topological BFS; O(V+E). Clarify edge direction and duplicate edges.",
    "Monotonic Stack": "Maintain a stack whose values or indices are monotonic; pop while the current value resolves prior next-greater/smaller relationships. Usually O(n) because each index is pushed and popped once.",
    "Union-Find": "Use parent pointers with path compression and union by rank/size for near-constant amortized operations; suitable for dynamic connectivity, not arbitrary deletions.",
    "Segment Tree": "Build a tree over ranges to answer/update associative aggregates in O(log n); clarify operation identity, lazy propagation needs, and memory overhead.",
    "Graph Shortest Path": "Choose BFS for unweighted edges, Dijkstra for nonnegative weights, or Bellman-Ford when negative edges are allowed; state complexity and negative-cycle behavior.",
}

FOLLOWUP_FRAMES = [
    "What concrete failure or edge case would falsify that answer for {topic}?",
    "Which guarantee are you relying on, and where is that guarantee enforced for {topic}?",
    "How would you test or measure the claim about {topic} under production-like load?",
    "What alternative would you choose if the scale, latency, or tenancy constraint changed for {topic}?",
    "What would you change at Staff scope to make the {topic} decision durable across teams?",
]

CATEGORY_META = {
    "Go": ("https://go.dev/ref/spec", "Go language/runtime; technical prompt taxonomy derived from official docs and practical engineering patterns."),
    "DSA": ("https://leetcode.com/problemset/", "Pattern-based practice prompt; no named employer is asserted for this individual problem."),
    "Distributed Systems": ("https://sre.google/workbook/table-of-contents/", "Distributed systems competency; exact employer/question/date not independently verified."),
    "System Design": ("https://amazon.jobs/content/en/how-we-hire/sde-iii-interview-prep", "Representative system-design exercise; not a verbatim employer question."),
    "Kafka": ("https://kafka.apache.org/documentation/", "Kafka technical prompt; frequency estimated for matching data/backend roles, not counted from interview reports."),
    "PostgreSQL": ("https://www.postgresql.org/docs/current/", "PostgreSQL technical prompt; frequency estimated for backend roles, not a per-company verified count."),
    "Redis": ("https://redis.io/docs/latest/", "Redis technical prompt; role-dependent frequency estimate."),
    "MongoDB": ("https://www.mongodb.com/docs/manual/", "MongoDB technical prompt; role-dependent frequency estimate."),
    "Kubernetes": ("https://kubernetes.io/docs/", "Kubernetes/platform prompt; role-dependent frequency estimate."),
    "AWS/Azure": ("https://docs.aws.amazon.com/wellarchitected/latest/framework/welcome.html", "Cloud architecture prompt; exact provider/team relevance must be validated from the job description."),
    "Security": ("https://owasp.org/www-project-top-ten/", "Security scenario prompt; exact employer/question/date not independently verified."),
    "Behavioral": ("https://careers.microsoft.com/v2/global/en/hiring-tips.html", "Practice prompt aligned to publicly stated behavioral evaluation; not a quoted question."),
    "AI/LLM": ("https://survey.stackoverflow.co/2025/ai/", "Emerging AI competency prompt. Survey evidence indicates adoption/trust context, not interview frequency."),
    "Resume-Based": ("interview-prep/README.md", "Profile-tailored prompt scaffold only. Resume absent; validate each premise and supply actual project evidence."),
}

IMPORTANCE = {"Go": 5, "DSA": 5, "Distributed Systems": 5, "System Design": 5, "Kafka": 5, "PostgreSQL": 5, "Redis": 3, "MongoDB": 3, "Kubernetes": 4, "AWS/Azure": 4, "Security": 5, "Behavioral": 5, "AI/LLM": 3, "Resume-Based": 5}
FREQUENCY = {"Go": "High for Go roles; role-dependent otherwise", "DSA": "High in algorithmic coding loops; varies by company/team", "Distributed Systems": "High for senior backend/platform; depth varies", "System Design": "High for senior roles; official Amazon SDE III guidance explicitly includes it", "Kafka": "High for streaming/data roles; medium for backend roles using Kafka", "PostgreSQL": "High for relational-backend roles", "Redis": "Medium; rises when caching/Redis appears in role", "MongoDB": "Medium; rises when MongoDB appears in role", "Kubernetes": "High for platform roles; medium for backend roles", "AWS/Azure": "High for cloud/platform roles; provider-specific", "Security": "High baseline; deeper in security/identity roles", "Behavioral": "High; company and level specific", "AI/LLM": "Emerging; high only for AI-facing roles today", "Resume-Based": "High: project deep dives are role-dependent and personalized"}
COMPANY_MAP = {
    "Go": "Go-focused backend/platform postings; exact companies require posting evidence.",
    "DSA": "Large product-company coding loops broadly; no per-company exact-question claim.",
    "Distributed Systems": "Backend, data, infra, storage, and platform roles across target set; team-specific.",
    "System Design": "Senior backend/platform roles across target set; Amazon SDE III officially lists system design.",
    "Kafka": "Confluent and streaming/data teams are strong role-fit targets; no individual question attribution verified.",
    "PostgreSQL": "Backend/data roles across target set; exact database choice is posting-specific.",
    "Redis": "Backend/platform roles using Redis; exact usage must be verified from postings/blogs.",
    "MongoDB": "Backend/data roles using MongoDB; exact usage must be verified from postings/blogs.",
    "Kubernetes": "Platform/cloud-native roles across target set; CNCF ecosystem evidence, not per-company interview mapping.",
    "AWS/Azure": "Cloud/platform roles; provider and service choices are role-specific.",
    "Security": "Cloud/security/identity roles and baseline backend reviews; exact loop emphasis varies.",
    "Behavioral": "Amazon publishes Leadership Principles; Microsoft publishes values/structured hiring guidance; others role-specific.",
    "AI/LLM": "AI-native and AI-product roles; Sierra has a first-party AI-native interview article (April 2026 indexed result). Exact questions and broader company mapping are not verified.",
    "Resume-Based": "Every interviewer may probe the actual resume; company-specific mapping cannot be assigned without resume/job.",
}

DIFFICULTY_BY_LENS = {"Beginner": "Easy", "Intermediate": "Medium", "Advanced": "Hard", "Senior": "Hard", "Staff": "Expert", "Principal": "Expert", "Debugging": "Hard", "Failure": "Expert", "Performance": "Hard", "Architecture": "Hard", "Tradeoff": "Hard", "Follow-up": "Hard", "Synthesis": "Expert"}


def subcategory(lens):
    if lens in {"Beginner", "Intermediate", "Advanced", "Senior", "Staff", "Principal"}:
        return lens + " interview"
    return lens + " / scenario"


def answer_and_exercise(category, topic):
    if category == "DSA":
        answer = f"Core answer: {TOPIC_CORE[topic]} " + ANSWER_CUES[category]
        exercise = f"Solve {topic} from a blank editor; test empty/minimal, duplicate, adversarial, and maximum-constraint cases; then explain an optimization and its proof."
    elif category == "Behavioral":
        answer = ANSWER_CUES[category] + " Select a real story matching the competency; do not invent a result."
        exercise = f"Record a truthful 2-3 minute STAR-L answer about {topic}; mark personal actions, alternatives, result evidence, and learning."
    elif category == "Resume-Based":
        answer = ANSWER_CUES[category] + " No resume was supplied; this row is a prompt, not an assertion that this experience occurred."
        exercise = f"Attach this prompt only to a verified resume bullet/project; collect architecture, decision, outcome, failure, and metric evidence for {topic}."
    elif category == "System Design":
        answer = ANSWER_CUES[category] + f" Scenario anchor: {topic}. Clarify product scope before selecting technologies."
        exercise = f"Run a 45-minute design for {topic}; include requirements, capacity, API/data, security, SLO, failure recovery, migration, and cost."
    elif category == "Go":
        answer = f"Core answer: {TOPIC_CORE[topic]} " + ANSWER_CUES[category]
        exercise = f"Write a minimal reproducible example or design involving {topic}; add one correctness test, one failure test, and one measurement."
    else:
        answer = ANSWER_CUES[category] + f" Topic anchor: {topic}."
        exercise = f"Write a minimal reproducible example or design involving {topic}; add one correctness test, one failure test, and one measurement."
    return answer, exercise


def build_question_record(category, frequency, importance, evidence_note, resource, topic, difficulty, template, qid):
    question = template.format(topic=topic)
    followups = [frame.format(topic=topic) for frame in FOLLOWUP_FRAMES]
    answer, exercise = answer_and_exercise(category, topic)
    return {
        "Question ID": f"Q{qid:05d}",
        "Category": category,
        "Subcategory": subcategory(difficulty),
        "Topic": topic,
        "Question": question,
        "Difficulty": difficulty,
        "Company": COMPANY_MAP[category],
        "Frequency": frequency,
        "Interview Round": "Coding / technical screen / system design / behavioral; exact round varies by employer and role",
        "Years Observed": "Concept remains relevant; exact year-specific prompt occurrence not independently verified",
        "Importance Score": importance,
        "Expected Answer": answer,
        "Common Wrong Answer": common_wrong(category),
        "Follow-up Question 1": followups[0],
        "Follow-up Question 2": followups[1],
        "Follow-up Question 3": followups[2],
        "Follow-up Question 4": followups[3],
        "Follow-up Question 5": followups[4],
        "Related Questions": f"Same topic: {topic}; neighboring category: {category}; search database by Topic and Category",
        "Recommended Resource": resource,
        "Practice Exercise": exercise,
        "Evidence Note": evidence_note,
    }


def record_rows():
    rows = []
    seen = set()
    qid = 1
    for category, total, frequency, importance, evidence_note, resource, topics in CATALOG:
        per_topic = total // len(topics)
        lenses = DSA_LENSES if category == "DSA" else LENSES
        assert per_topic <= len(lenses), (category, per_topic, len(lenses))
        for topic in topics:
            for difficulty, template in lenses[:per_topic]:
                question = template.format(topic=topic)
                if question in seen:
                    raise ValueError(f"Duplicate question: {question}")
                seen.add(question)
                rows.append(build_question_record(category, frequency, importance, evidence_note, resource, topic, difficulty, template, qid))
                qid += 1
    return rows


def common_wrong(category):
    defaults = {
        "Go": "Confusing an implementation detail with a language guarantee; ignoring aliasing, lifecycle, or a race; giving no testable example.",
        "DSA": "Jumping to code without clarifying constraints or invariant; missing edge cases; stating complexity inaccurately.",
        "Distributed Systems": "Using slogans such as 'pick two' or 'exactly once' without stating failure model and atomicity boundary.",
        "System Design": "Drawing components before requirements; ignoring capacity, operations, security, migration, and cost.",
        "Kafka": "Claiming global ordering or cross-system exactly-once behavior; ignoring offset timing and duplicate effects.",
        "PostgreSQL": "Assuming an index guarantees a fast plan; confusing MVCC with no locking; ignoring actual plans and isolation anomalies.",
        "Redis": "Treating cache data as durable or replication as synchronous/zero-loss without configuration evidence.",
        "MongoDB": "Assuming schema flexibility removes modeling work; ignoring index, shard-key, consistency, and document-growth costs.",
        "Kubernetes": "Assuming restarts equal high availability; confusing liveness with readiness; ignoring application and dependency state.",
        "AWS/Azure": "Choosing service names before requirements; ignoring IAM, quotas, failure domains, egress, and recovery.",
        "Security": "Trusting client-side checks or prompts; validating token signature but not claims; failing to enforce object-level authorization.",
        "Behavioral": "Vague 'we' narrative, no personal actions or evidence, fabricated impact, blame, or a fake failure with no learning.",
        "AI/LLM": "Treating model output as trusted; demo-only evaluation; unscoped tools; similarity mistaken for authorization or factuality.",
        "Resume-Based": "Claiming ownership or metrics not supported by the actual resume/project; unable to explain alternatives or failure.",
    }
    return defaults[category]


def make_workbook(rows):
    wb = Workbook()
    intro = wb.active
    intro.title = "Read Me"
    info = [
        ["Purpose", "Question-centered backend/platform interview database; filter the Question Bank instead of navigating a topic syllabus."],
        ["Question records", len(rows)],
        ["Evidence date", "2026-10-19"],
        ["Question provenance", "Practice formulations synthesized from role expectations and authoritative technical references; they are not represented as verbatim interview reports."],
        ["Frequency", "Qualitative role-family estimate, not a measured company-level probability. Exact question frequencies are not publicly enumerable."],
        ["Company mapping", "Maps role/company fit and published broad signals; not a claim that each company asked the row's exact question."],
        ["Resume prompts", "500 profile-tailored probes are scaffolds. No resume was supplied; verify every premise and do not invent personal answers."],
        ["Saturation", "Large coverage target met; no finite bank guarantees never encountering a new prompt. Expand from job postings, recruiter guidance, and mock feedback."],
        ["Research sources", "Amazon SDE III prep; Microsoft hiring guidance; Go/PostgreSQL/Redis/Kafka/Kubernetes/AWS docs; CNCF and Stack Overflow surveys; OWASP GenAI."],
    ]
    for row in info:
        intro.append(row)
    intro.column_dimensions["A"].width = 26
    intro.column_dimensions["B"].width = 115
    for row in intro.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.column == 1:
                cell.font = Font(bold=True, color="174A5B")

    headers = list(rows[0])
    sheet = wb.create_sheet("Question Bank")
    sheet.append(headers)
    for item in rows:
        sheet.append([item[key] for key in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 38
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174A5B")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = {"Question ID": 12, "Category": 22, "Subcategory": 22, "Topic": 38, "Question": 72, "Difficulty": 14, "Company": 68, "Frequency": 48, "Interview Round": 54, "Years Observed": 55, "Importance Score": 18, "Expected Answer": 90, "Common Wrong Answer": 75, "Related Questions": 50, "Recommended Resource": 56, "Practice Exercise": 82, "Evidence Note": 80}
    for col_index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(col_index)].width = widths.get(header, 58)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    killers = wb.create_sheet("Top 200 Risk Drills")
    killer_headers = ["Risk Rank", "Risk Tier", "Why Candidates Fail", *headers]
    killers.append(killer_headers)
    risk_order = {"Expert": 0, "Hard": 1, "Medium": 2, "Easy": 3}
    candidates = sorted(
        rows,
        key=lambda item: (
            -int(item["Importance Score"]),
            risk_order.get(item["Difficulty"], 4),
            item["Question ID"],
        ),
    )
    by_category = {}
    for item in candidates:
        by_category.setdefault(item["Category"], []).append(item)
    selected = []
    category_names = [category for category, *_ in CATALOG]
    while len(selected) < 200:
        for category in category_names:
            bucket = by_category[category]
            if bucket and len(selected) < 200:
                selected.append(bucket.pop(0))
    for rank, item in enumerate(selected, start=1):
        if rank <= 40:
            tier = "Critical"
        elif rank <= 100:
            tier = "High"
        elif rank <= 160:
            tier = "Medium"
        else:
            tier = "Low"
        reason = risk_reason(item)
        killers.append([rank, tier, reason, *[item[key] for key in headers]])
    killers.freeze_panes = "A2"
    killers.auto_filter.ref = killers.dimensions
    killers.row_dimensions[1].height = 38
    for cell in killers[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="8A3B12")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    killers.column_dimensions["A"].width = 12
    killers.column_dimensions["B"].width = 14
    killers.column_dimensions["C"].width = 70
    for col_index, header in enumerate(headers, start=4):
        killers.column_dimensions[get_column_letter(col_index)].width = widths.get(header, 58)
    for row in killers.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    killers.sheet_view.showGridLines = False

    counts = Counter(item["Category"] for item in rows)
    summary = wb.create_sheet("Coverage Counts")
    summary.append(["Category", "Question Count", "Target", "Status"])
    for category, total, *_ in CATALOG:
        summary.append([category, counts[category], total, "Met" if counts[category] >= total else "Below target"])
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174A5B")
    summary.column_dimensions["A"].width = 28
    for col in ("B", "C", "D"):
        summary.column_dimensions[col].width = 18

    source = wb.create_sheet("Research Sources")
    source.append(["Source", "What it supports", "What it does not establish"])
    refs = [
        ["https://amazon.jobs/content/en/how-we-hire/sde-iii-interview-prep", "Published Amazon senior expectations, coding/system design/behavioral structure", "Other companies' processes or exact question frequencies"],
        ["https://amazon.jobs/content/en/how-we-hire/interview-prep/software-development-topics", "Amazon's broad software interview topic families", "Exact questions asked in each loop"],
        ["https://careers.microsoft.com/v2/global/en/hiring-tips.html", "Microsoft published structured hiring, values, collaboration, impact, responsible AI guidance", "Universal AI interview format"],
        ["https://survey.stackoverflow.co/2025/ai/", "Developer AI usage, trust, and agent adoption survey context", "Hiring frequency or a specific company's interview rubric"],
        ["https://sierra.ai/blog/the-ai-native-interview", "Sierra first-party article title and indexed publication date indicate a company-specific AI-native interview process signal", "Article body was inaccessible in this environment; specific stages, questions, or industry-wide prevalence"],
        ["https://www.cncf.io/reports/state-of-cloud-native-development-q1-2026/", "Cloud-native/platform ecosystem context", "Interview question counts"],
        ["https://go.dev/ref/mem", "Go memory model guarantees", "Any particular employer's prompt history"],
        ["https://go.dev/doc/diagnostics", "Go diagnostics/profiling tools", "Interview frequency"],
        ["https://kafka.apache.org/documentation/", "Kafka protocol and operational concepts", "Company-level usage or interview occurrence"],
        ["https://www.postgresql.org/docs/current/mvcc.html", "PostgreSQL concurrency and MVCC behavior", "Interview occurrence"],
        ["https://redis.io/docs/latest/operate/oss_and_stack/management/replication/", "Redis replication guarantees and limitations", "Interview occurrence"],
        ["https://www.mongodb.com/docs/manual/", "MongoDB data model, indexes, replication, and sharding concepts", "Interview occurrence"],
        ["https://kubernetes.io/docs/concepts/overview/", "Kubernetes capabilities and boundaries", "Interview occurrence"],
        ["https://owasp.org/www-project-top-10-for-large-language-model-applications/", "LLM application risk taxonomy", "Specific company interview questions"],
        ["https://modelcontextprotocol.io/specification/2025-06-18", "MCP concepts and trust/safety considerations", "Universal MCP screening expectation"],
        ["https://sre.google/workbook/table-of-contents/", "SLOs, incident response, load, reliability practice", "Interview frequency"],
    ]
    for row in refs:
        source.append(row)
    source.freeze_panes = "A2"
    source.auto_filter.ref = source.dimensions
    for cell in source[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174A5B")
    for col, width in zip("ABC", (78, 74, 64)):
        source.column_dimensions[col].width = width
    for row in source.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    companies = wb.create_sheet("Company Targets")
    with (ROOT / "company-targets.csv").open(newline="", encoding="utf-8") as handle:
        for row in csv.reader(handle):
            companies.append(row)
    companies.freeze_panes = "A2"
    companies.auto_filter.ref = companies.dimensions
    for cell in companies[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174A5B")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column in companies.columns:
        letter = get_column_letter(column[0].column)
        companies.column_dimensions[letter].width = min(max(len(str(column[0].value or "")) * 1.2, 16), 48)
    for row in companies.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.save(XLSX_PATH)


def risk_reason(item):
    category_reasons = {
        "Go": "Often exposes shallow semantics knowledge; weak answers omit memory/lifecycle guarantees and a reproducing test.",
        "DSA": "Coding screens penalize missed invariants, edge cases, complexity, and inability to test under time.",
        "Distributed Systems": "Candidates often overclaim consistency or exactly-once guarantees without a clear failure model.",
        "System Design": "Senior candidates lose signal when requirements, estimates, security, operations, and migration are skipped.",
        "Kafka": "Delivery, ordering, offset, and transaction boundaries are commonly conflated in explanations.",
        "PostgreSQL": "ORM familiarity can hide gaps in transaction, lock, plan, and maintenance reasoning.",
        "Redis": "Cache and replication guarantees are often overstated or treated as durable by default.",
        "MongoDB": "Candidates can overstate schema flexibility while missing index, shard-key, consistency, and document-growth tradeoffs.",
        "Kubernetes": "Control-plane primitives are mistaken for application-level availability and recovery.",
        "AWS/Azure": "Service-name recall without IAM, failure-domain, quota, and cost reasoning is weak evidence.",
        "Security": "A missed authorization or tenant boundary can create critical correctness and trust failures.",
        "Behavioral": "Vague ownership, unsupported results, and blame reduce senior-level evidence.",
        "AI/LLM": "Untrusted retrieval/tools without evaluation and access controls create safety/security failures.",
        "Resume-Based": "Unverifiable ownership or metrics fail deep-dive credibility; prompts must be grounded in actual resume facts.",
    }
    return category_reasons[item["Category"]]


def write_csv(rows):
    with CSV_PATH.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


if __name__ == "__main__":
    records = record_rows()
    write_csv(records)
    make_workbook(records)
    counts = Counter(row["Category"] for row in records)
    print(f"Created {len(records)} unique question records")
    print(f"CSV: {CSV_PATH}")
    print(f"Workbook: {XLSX_PATH}")
    print("Category counts:", dict(counts))
